import os
import requests
from io import BytesIO
import subprocess
import sys

def install_and_import():
    try:
        import docx
    except ImportError:
        print("Installing python-docx...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "requests"])
        
install_and_import()

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_document():
    # Download the logo
    logo_url = "https://unigrande.edu.br/_next/static/media/logo-unigrande-color-blue.0251e5cb.png"
    try:
        response = requests.get(logo_url)
        response.raise_for_status()
        logo_image = BytesIO(response.content)
    except Exception as e:
        print(f"Error downloading logo: {e}")
        logo_image = None

    document = Document()

    # Configure ABNT margins
    sections = document.sections
    for section in sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)

    # Base font style
    style = document.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # ===== CAPA (Cover) =====
    
    # Institution / Logo
    if logo_image:
        p_logo = document.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_image, width=Inches(2.5))
    
    p_inst = document.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_inst.add_run("CENTRO UNIVERSITÁRIO UNIGRANDE")
    run_inst.bold = True

    # Spacing
    for _ in range(5):
        document.add_paragraph()

    # Author
    p_author = document.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_author = p_author.add_run("ANTONIO ERICK CONCEIÇÃO DA SILVA")
    run_author.bold = True

    # Spacing
    for _ in range(6):
        document.add_paragraph()

    # Title
    p_title = document.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("6 ARTICULAÇÕES SINOVIAIS")
    run_title.bold = True

    # Spacing to bottom
    for _ in range(8):
        document.add_paragraph()

    # Location / Year
    p_loc = document.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc.add_run("Buriticupu - MA\n2026").bold = True

    document.add_page_break()

    # ===== CONTRA CAPA (Inside Cover) =====
    
    # Author
    p_author2 = document.add_paragraph()
    p_author2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_author2 = p_author2.add_run("ANTONIO ERICK CONCEIÇÃO DA SILVA")
    run_author2.bold = True

    # Spacing
    for _ in range(6):
        document.add_paragraph()

    # Title
    p_title2 = document.add_paragraph()
    p_title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title2 = p_title2.add_run("6 ARTICULAÇÕES SINOVIAIS")
    run_title2.bold = True

    for _ in range(2):
        document.add_paragraph()

    # Note
    p_note = document.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Move note to the right (ABNT style)
    p_note.paragraph_format.left_indent = Cm(8)
    run_note = p_note.add_run(
        "Trabalho acadêmico apresentado ao Centro Universitário UniGrande, "
        "como requisito parcial para avaliação.\n\n"
        "Professora: Bianca Lira"
    )

    # Spacing to bottom
    for _ in range(9):
        document.add_paragraph()

    # Location / Year
    p_loc2 = document.add_paragraph()
    p_loc2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc2.add_run("Buriticupu - MA\n2026").bold = True

    document.add_page_break()

    # ===== TEXTO (Content) =====
    
    # Title
    p_content_title = document.add_paragraph()
    p_content_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_content_title.add_run("6 ARTICULAÇÕES SINOVIAIS").bold = True
    
    document.add_paragraph() # space

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.first_line_indent = Cm(1.25)
    intro.add_run(
        "As articulações sinoviais são o tipo mais comum e mais móvel de articulação no corpo humano. "
        "Elas se caracterizam pela presença de uma cápsula articular que envolve uma cavidade sinovial cheia de líquido sinovial, "
        "o qual atua como lubrificante. As superfícies ósseas em contato são recobertas por cartilagem hialina (cartilagem articular), "
        "garantindo um movimento suave e sem atrito. A seguir, são descritas 6 importantes articulações sinoviais do corpo humano."
    )

    # Joints data
    joints = [
        {
            "name": "1. Articulação Glenoumeral (Ombro)",
            "type": "Sinovial Esferóide",
            "bones": "Escápula (cavidade glenóide) e Úmero (cabeça do úmero).",
            "desc": "É uma articulação multiaxial, permitindo uma ampla gama de movimentos. Devido à sua grande mobilidade, é também uma das articulações mais instáveis do corpo, dependendo fortemente do manguito rotador e dos ligamentos associados para sua estabilização.",
            "movements": "Flexão, extensão, abdução, adução, rotação medial, rotação lateral e circundução."
        },
        {
            "name": "2. Articulação Coxofemoral (Quadril)",
            "type": "Sinovial Esferóide",
            "bones": "Osso do quadril (acetábulo) e Fêmur (cabeça do fêmur).",
            "desc": "Semelhante à articulação do ombro por ser esferóide, mas projetada para sustentação de peso e estabilidade profunda. O acetábulo é profundo e engloba grande parte da cabeça femoral, o que reduz sua amplitude em relação ao ombro, mas fornece uma articulação incrivelmente forte.",
            "movements": "Flexão, extensão, abdução, adução, rotação medial, rotação lateral e circundução."
        },
        {
            "name": "3. Articulação do Cotovelo (Umeroulnar e Umerorradial)",
            "type": "Sinovial Gínglimo (Dobradiça)",
            "bones": "Úmero (tróclea e capítulo), Ulna (incisura troclear) e Rádio (cabeça do rádio).",
            "desc": "Trata-se de uma articulação complexa que funciona principalmente como uma dobradiça (gínglimo), restringindo o movimento em grande parte a um único plano. É estabilizada por fortes ligamentos colaterais (ulnar e radial).",
            "movements": "Flexão e extensão."
        },
        {
            "name": "4. Articulação Tibiofemoral (Joelho)",
            "type": "Sinovial Gínglimo Modificado (ou Bicondilar)",
            "bones": "Fêmur (côndilos femorais) e Tíbia (côndilos tibiais), além da Patela (articulação patelofemoral).",
            "desc": "É a maior articulação do corpo humano e suporta grande parte do peso corporal. Devido à incongruência entre os ossos, apresenta os meniscos (estruturas fibrocartilaginosas) que melhoram o encaixe. Permite primariamente flexão e extensão, com uma leve rotação quando flexionada.",
            "movements": "Flexão, extensão e leve rotação medial/lateral (quando flexionado)."
        },
        {
            "name": "5. Articulação Radiocarpal (Punho)",
            "type": "Sinovial Elipsóide (Condilar)",
            "bones": "Rádio (extremidade distal) e ossos do Carpo (escafoide, semilunar e piramidal).",
            "desc": "Esta articulação conecta o antebraço à mão. Sendo condilar, permite movimento em dois planos principais. O disco articular a separa da ulna, que não participa diretamente dessa articulação.",
            "movements": "Flexão, extensão, abdução (desvio radial), adução (desvio ulnar) e circundução."
        },
        {
            "name": "6. Articulação Temporomandibular (ATM)",
            "type": "Sinovial Gínglimo Modificado / Plana",
            "bones": "Osso Temporal (fossa mandibular e tubérculo articular) e Mandíbula (cabeça ou côndilo da mandíbula).",
            "desc": "É uma das articulações mais utilizadas no corpo humano (mastigação, fala, deglutição). Possui um disco articular que divide a cavidade sinovial em duas partes, permitindo tanto um movimento de dobradiça quanto um movimento de deslizamento.",
            "movements": "Elevação, depressão (abertura e fechamento da boca), protrusão, retração e movimentos de lateralidade."
        }
    ]

    for joint in joints:
        document.add_paragraph()
        
        p_name = document.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run_name = p_name.add_run(joint['name'])
        run_name.bold = True
        
        p_det = document.add_paragraph()
        p_det.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_det.paragraph_format.first_line_indent = Cm(1.25)
        
        # Build the text with bold labels
        run_class = p_det.add_run("Classificação:")
        run_class.bold = True
        p_det.add_run(f" {joint['type']}. ")
        
        run_bones = p_det.add_run("Ossos Envolvidos:")
        run_bones.bold = True
        p_det.add_run(f" {joint['bones']} ")
        
        run_mov = p_det.add_run("Movimentos:")
        run_mov.bold = True
        p_det.add_run(f" {joint['movements']}\n")
        
        p_desc = document.add_paragraph()
        p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_desc.add_run(joint['desc'])

    output_path = os.path.join(os.getcwd(), "Trabalho_Articulacoes_Sinoviais.docx")
    document.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == "__main__":
    create_document()
